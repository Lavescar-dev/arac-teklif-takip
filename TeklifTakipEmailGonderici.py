#!/usr/bin/env python3
"""
Quote tracking reminder sender.

Rules:
- Day 2 and still pending: notify common email
- Day 4 and still pending: notify boss
"""

from __future__ import annotations

import argparse
import logging
import re
import smtplib
import sys
import traceback
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime
from email.message import EmailMessage
from pathlib import Path
from typing import Dict, List, Optional, Sequence

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel


DEFAULT_CONFIG = "TeklifTakipAyarlar.txt"
DEFAULT_LOG = "TeklifTakip.log"


def base_dir() -> Path:
    """Folder where config/log/excel live. Works for both .py and PyInstaller .exe."""
    if getattr(sys, "frozen", False):
        return Path(sys.executable).resolve().parent
    return Path(__file__).resolve().parent


def ensure_config(script_dir: Path, config_name: str, force_setup: bool) -> bool:
    config_path = script_dir / config_name
    if not force_setup and config_path.exists():
        return True

    try:
        from setup_wizard import run_teklif_takip_wizard
    except ImportError:
        if not config_path.exists():
            logging.error(
                "Config dosyasi yok ve setup wizard yuklenemedi. %s olusturup tekrar deneyin.",
                config_path,
            )
            return False
        return True

    return run_teklif_takip_wizard(config_path)


DATE_ALIASES = ("TARIH", "TARIHI", "KAYITTARIHI")
STATUS_ALIASES = ("SONDURUMU", "DURUM", "TEKLIFDURUMU", "SONDURUM")
COMPANY_ALIASES = ("SIRKETADI", "FIRMAADI", "MUSTERIADI", "MUSTERI", "FIRMA")
CONTACT_ALIASES = ("ILGILI", "YETKILI", "SORUMLU")
FLAG_COMMON_ALIASES = ("MAIL1", "REMINDER_3_SENT", "REMINDER3SENT", "MAIL_1")
FLAG_BOSS_ALIASES = ("MAIL2", "REMINDER_5_SENT", "REMINDER5SENT", "MAIL_2")
TURKISH_MONTHS = {
    "OCAK": 1,
    "SUBAT": 2,
    "MART": 3,
    "NISAN": 4,
    "MAYIS": 5,
    "HAZIRAN": 6,
    "TEMMUZ": 7,
    "AGUSTOS": 8,
    "EYLUL": 9,
    "EKIM": 10,
    "KASIM": 11,
    "ARALIK": 12,
}


@dataclass
class Config:
    filename: str = ""
    sheetname: str = ""
    smtp_host: str = "smtp.gmail.com"
    smtp_port: int = 465
    smtp_username: str = ""
    smtp_password: str = ""
    common_email: str = ""
    boss_email: str = ""
    cc_emails: List[str] = field(default_factory=list)
    reminder_common_days: int = 2
    reminder_boss_days: int = 4
    date_column: str = "TARIH"
    status_column: str = "SON DURUMU"
    company_column: str = "SIRKET ADI"
    contact_column: str = "ILGILI"
    flag_common_column: str = "MAIL1"
    flag_common_at_column: str = ""
    flag_boss_column: str = "MAIL2"
    flag_boss_at_column: str = ""
    subject_prefix: str = "[Teklif Takip]"


@dataclass
class PendingRecord:
    row_number: int
    record_date: date
    days_passed: int
    company: str
    contact: str
    status: str


def configure_logging(log_file: Path) -> None:
    logging.basicConfig(
        filename=str(log_file),
        level=logging.INFO,
        format="%(asctime)s - %(levelname)s - %(message)s",
    )


def read_text_with_fallback(path: Path) -> List[str]:
    encodings = ("utf-8-sig", "utf-8", "cp1254", "latin-1")
    for encoding in encodings:
        try:
            return path.read_text(encoding=encoding).splitlines()
        except UnicodeDecodeError:
            continue
    raise UnicodeDecodeError("unknown", b"", 0, 1, "Config decode failed")


def parse_int(raw: str, default_value: int) -> int:
    try:
        return int(raw.strip())
    except Exception:
        return default_value


def parse_email_list(raw: str) -> List[str]:
    if not raw:
        return []
    parts = re.split(r"[;,]", raw)
    return [p.strip() for p in parts if p.strip()]


def parse_config(path: Path) -> Config:
    cfg = Config()
    lines = read_text_with_fallback(path)

    for raw_line in lines:
        line = raw_line.strip()
        if not line or line.startswith("#"):
            continue
        if ":" not in line:
            continue
        key, value = line.split(":", 1)
        key = key.strip().upper()
        value = value.strip()

        if key == "FILENAME":
            cfg.filename = value
        elif key == "SHEETNAME":
            cfg.sheetname = value
        elif key == "SMTP_HOST":
            cfg.smtp_host = value or cfg.smtp_host
        elif key == "SMTP_PORT":
            cfg.smtp_port = parse_int(value, cfg.smtp_port)
        elif key == "SMTP_USERNAME":
            cfg.smtp_username = value
        elif key == "SMTP_PASSWORD":
            cfg.smtp_password = value
        elif key == "COMMON_EMAIL":
            cfg.common_email = value
        elif key == "EMPLOYEE_EMAIL":
            # backward compatible alias
            cfg.common_email = value
        elif key == "BOSS_EMAIL":
            cfg.boss_email = value
        elif key == "CC_EMAILS":
            cfg.cc_emails = [x.strip() for x in value.split(",") if x.strip()]
        elif key == "REMINDER_COMMON_DAYS":
            cfg.reminder_common_days = parse_int(value, cfg.reminder_common_days)
        elif key == "REMINDER_BOSS_DAYS":
            cfg.reminder_boss_days = parse_int(value, cfg.reminder_boss_days)
        elif key == "REMINDER_3_DAYS":
            # backward compatible alias
            cfg.reminder_common_days = parse_int(value, cfg.reminder_common_days)
        elif key == "REMINDER_5_DAYS":
            # backward compatible alias
            cfg.reminder_boss_days = parse_int(value, cfg.reminder_boss_days)
        elif key == "DATE_COLUMN":
            cfg.date_column = value or cfg.date_column
        elif key == "STATUS_COLUMN":
            cfg.status_column = value or cfg.status_column
        elif key == "COMPANY_COLUMN":
            cfg.company_column = value or cfg.company_column
        elif key == "CONTACT_COLUMN":
            cfg.contact_column = value or cfg.contact_column
        elif key == "FLAG_COMMON_COLUMN":
            cfg.flag_common_column = value or cfg.flag_common_column
        elif key == "FLAG_COMMON_AT_COLUMN":
            cfg.flag_common_at_column = value
        elif key == "FLAG_BOSS_COLUMN":
            cfg.flag_boss_column = value or cfg.flag_boss_column
        elif key == "FLAG_BOSS_AT_COLUMN":
            cfg.flag_boss_at_column = value
        elif key == "FLAG_3_COLUMN":
            # backward compatible alias
            cfg.flag_common_column = value or cfg.flag_common_column
        elif key == "FLAG_3_AT_COLUMN":
            cfg.flag_common_at_column = value
        elif key == "FLAG_5_COLUMN":
            # backward compatible alias
            cfg.flag_boss_column = value or cfg.flag_boss_column
        elif key == "FLAG_5_AT_COLUMN":
            cfg.flag_boss_at_column = value
        elif key == "SUBJECT_PREFIX":
            cfg.subject_prefix = value or cfg.subject_prefix

    required = (
        ("FILENAME", cfg.filename),
        ("SMTP_USERNAME", cfg.smtp_username),
        ("SMTP_PASSWORD", cfg.smtp_password),
        ("COMMON_EMAIL", cfg.common_email),
        ("BOSS_EMAIL", cfg.boss_email),
    )
    missing = [k for k, v in required if not v]
    if missing:
        raise ValueError(f"Missing required config keys: {', '.join(missing)}")

    if cfg.reminder_boss_days < cfg.reminder_common_days:
        raise ValueError("REMINDER_BOSS_DAYS cannot be smaller than REMINDER_COMMON_DAYS")

    return cfg


def normalize_text(value: object) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if not text:
        return ""
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = text.upper()
    text = re.sub(r"[^A-Z0-9]+", "", text)
    return text


def parse_excel_date(value: object) -> Optional[date]:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        try:
            return from_excel(value).date()
        except Exception:
            return None

    text = str(value).strip()
    if not text:
        return None

    date_formats = (
        "%d.%m.%Y",
        "%d/%m/%Y",
        "%Y-%m-%d",
        "%d-%m-%Y",
        "%d.%m.%Y %H:%M",
        "%d/%m/%Y %H:%M",
    )
    for fmt in date_formats:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue

    # Turkish textual date support, e.g. "16 Ocak 2026 Cuma"
    normalized = normalize_text(text)
    match = re.match(r"^(\d{1,2})([A-Z]+)(\d{4})", normalized)
    if match:
        day_raw, month_raw, year_raw = match.groups()
        month = TURKISH_MONTHS.get(month_raw)
        if month:
            try:
                return date(int(year_raw), month, int(day_raw))
            except ValueError:
                return None
    return None


def is_true_flag(value: object) -> bool:
    text = normalize_text(value)
    return text in {"1", "TRUE", "YES", "EVET", "GONDERILDI"}


def is_pending_status(value: object) -> bool:
    text = normalize_text(value)
    if not text:
        return True
    closed_tokens = ("VERILDI", "KAPANDI", "IPTAL", "VAZGECILDI", "ONAYLANDI")
    return not any(token in text for token in closed_tokens)


def build_header_map(sheet) -> Dict[str, int]:
    header_map: Dict[str, int] = {}
    for col_idx in range(1, sheet.max_column + 1):
        header = sheet.cell(row=1, column=col_idx).value
        key = normalize_text(header)
        if key and key not in header_map:
            header_map[key] = col_idx
    return header_map


def find_column(header_map: Dict[str, int], preferred_name: str, aliases: Sequence[str]) -> Optional[int]:
    preferred_key = normalize_text(preferred_name)
    if preferred_key in header_map:
        return header_map[preferred_key]

    for alias in aliases:
        alias_key = normalize_text(alias)
        if alias_key in header_map:
            return header_map[alias_key]
    return None


def ensure_column(sheet, header_map: Dict[str, int], column_name: str) -> int:
    existing = find_column(header_map, column_name, ())
    if existing is not None:
        return existing

    col_idx = sheet.max_column + 1
    sheet.cell(row=1, column=col_idx).value = column_name
    header_map[normalize_text(column_name)] = col_idx
    return col_idx


def format_record_line(record: PendingRecord) -> str:
    return (
        f"- Satir {record.row_number}: Sirket={record.company or '-'} | "
        f"Ilgili={record.contact or '-'} | Tarih={record.record_date:%d.%m.%Y} | "
        f"GecenGun={record.days_passed} | Durum={record.status or '-'}"
    )


def build_email(subject: str, records: List[PendingRecord]) -> str:
    lines = [
        "Merhaba,",
        "",
        "Asagidaki teklif kayitlari icin hatirlatma olustu:",
        "",
    ]
    lines.extend(format_record_line(record) for record in records)
    lines.extend(["", "Bu e-posta otomatik gonderilmistir."])
    return "\n".join(lines)


def send_email(
    cfg: Config,
    to_email: str,
    subject: str,
    body: str,
    extra_cc: Optional[List[str]] = None,
) -> None:
    primary_recipients = parse_email_list(to_email)
    if not primary_recipients:
        raise ValueError("No recipient configured for this reminder.")

    msg = EmailMessage()
    msg["Subject"] = subject
    msg["From"] = cfg.smtp_username
    msg["To"] = ", ".join(primary_recipients)

    cc_list: List[str] = list(cfg.cc_emails)
    if extra_cc:
        cc_list.extend(extra_cc)
    cc_list = [x for x in dict.fromkeys(cc_list) if x]
    if cc_list:
        msg["Cc"] = ", ".join(cc_list)

    msg.set_content(body)

    all_recipients = primary_recipients + cc_list
    with smtplib.SMTP_SSL(cfg.smtp_host, cfg.smtp_port, timeout=30) as smtp:
        smtp.login(cfg.smtp_username, cfg.smtp_password)
        smtp.send_message(msg, to_addrs=all_recipients)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--config", default=DEFAULT_CONFIG, help="Config file path")
    parser.add_argument("--dry-run", action="store_true", help="No email, no write")
    parser.add_argument("--setup", action="store_true", help="Open the Tkinter setup wizard.")
    args = parser.parse_args()

    script_dir = base_dir()
    config_path = script_dir / args.config
    log_path = script_dir / DEFAULT_LOG
    configure_logging(log_path)

    if not ensure_config(script_dir, args.config, force_setup=args.setup):
        return 1

    if not config_path.exists():
        logging.error("Config file not found: %s", config_path)
        return 1

    try:
        cfg = parse_config(config_path)
        logging.info("Config parsed successfully.")
    except Exception:
        logging.exception("Config parse failed.")
        logging.error(traceback.format_exc())
        return 1

    excel_path = script_dir / cfg.filename
    if not excel_path.exists():
        logging.error("Excel file not found: %s", excel_path)
        return 1

    wb = None
    try:
        wb = load_workbook(excel_path)
        sheet = wb[cfg.sheetname] if cfg.sheetname and cfg.sheetname in wb.sheetnames else wb.worksheets[0]
        header_map = build_header_map(sheet)

        date_col = find_column(header_map, cfg.date_column, DATE_ALIASES)
        status_col = find_column(header_map, cfg.status_column, STATUS_ALIASES)
        company_col = find_column(header_map, cfg.company_column, COMPANY_ALIASES)
        contact_col = find_column(header_map, cfg.contact_column, CONTACT_ALIASES)

        if date_col is None or status_col is None or company_col is None:
            logging.error(
                "Required column missing. date=%s status=%s company=%s",
                date_col,
                status_col,
                company_col,
            )
            return 1

        if args.dry_run:
            flag_common_col = find_column(header_map, cfg.flag_common_column, FLAG_COMMON_ALIASES)
            flag_common_at_col = (
                find_column(header_map, cfg.flag_common_at_column, ()) if cfg.flag_common_at_column else None
            )
            flag_boss_col = find_column(header_map, cfg.flag_boss_column, FLAG_BOSS_ALIASES)
            flag_boss_at_col = (
                find_column(header_map, cfg.flag_boss_at_column, ()) if cfg.flag_boss_at_column else None
            )
        else:
            existing_common = find_column(header_map, cfg.flag_common_column, FLAG_COMMON_ALIASES)
            flag_common_col = existing_common or ensure_column(sheet, header_map, cfg.flag_common_column)

            existing_boss = find_column(header_map, cfg.flag_boss_column, FLAG_BOSS_ALIASES)
            flag_boss_col = existing_boss or ensure_column(sheet, header_map, cfg.flag_boss_column)

            flag_common_at_col = (
                ensure_column(sheet, header_map, cfg.flag_common_at_column) if cfg.flag_common_at_column else None
            )
            flag_boss_at_col = (
                ensure_column(sheet, header_map, cfg.flag_boss_at_column) if cfg.flag_boss_at_column else None
            )

        today = date.today()
        common_records: List[PendingRecord] = []
        boss_records: List[PendingRecord] = []

        for row in range(2, sheet.max_row + 1):
            raw_date = sheet.cell(row=row, column=date_col).value
            record_date = parse_excel_date(raw_date)
            if record_date is None:
                continue

            days_passed = (today - record_date).days
            if days_passed < 0:
                continue

            status_value = sheet.cell(row=row, column=status_col).value
            if not is_pending_status(status_value):
                continue

            company = str(sheet.cell(row=row, column=company_col).value or "").strip()
            contact = ""
            if contact_col is not None:
                contact = str(sheet.cell(row=row, column=contact_col).value or "").strip()
            status = str(status_value or "").strip()

            sent_common = False
            sent_boss = False
            if flag_common_col is not None:
                sent_common = is_true_flag(sheet.cell(row=row, column=flag_common_col).value)
            if flag_boss_col is not None:
                sent_boss = is_true_flag(sheet.cell(row=row, column=flag_boss_col).value)

            rec = PendingRecord(
                row_number=row,
                record_date=record_date,
                days_passed=days_passed,
                company=company,
                contact=contact,
                status=status,
            )

            if days_passed >= cfg.reminder_common_days and not sent_common:
                common_records.append(rec)
            if days_passed >= cfg.reminder_boss_days and not sent_boss:
                boss_records.append(rec)

        if args.dry_run:
            print(f"DRY RUN: common reminders={len(common_records)} boss reminders={len(boss_records)}")
            for rec in common_records[:5]:
                print("ORTAK", format_record_line(rec))
            for rec in boss_records[:5]:
                print("BOS", format_record_line(rec))
            logging.info(
                "Dry run completed. common=%s boss=%s",
                len(common_records),
                len(boss_records),
            )
            return 0

        now_text = datetime.now().strftime("%Y-%m-%d %H:%M")
        changed = False

        if common_records:
            subject = (
                f"{cfg.subject_prefix} {len(common_records)} kayit "
                f"{cfg.reminder_common_days} gunu asti (ortak mail)"
            )
            body = build_email(subject, common_records)
            send_email(cfg, cfg.common_email, subject, body)
            for rec in common_records:
                sheet.cell(row=rec.row_number, column=flag_common_col).value = "YES"
                if flag_common_at_col is not None:
                    sheet.cell(row=rec.row_number, column=flag_common_at_col).value = now_text
            changed = True
            logging.info("Common reminder sent for %s records.", len(common_records))

        if boss_records:
            subject = (
                f"{cfg.subject_prefix} {len(boss_records)} kayit "
                f"{cfg.reminder_boss_days} gunu asti (patron)"
            )
            body = build_email(subject, boss_records)
            send_email(cfg, cfg.boss_email, subject, body)
            for rec in boss_records:
                sheet.cell(row=rec.row_number, column=flag_boss_col).value = "YES"
                if flag_boss_at_col is not None:
                    sheet.cell(row=rec.row_number, column=flag_boss_at_col).value = now_text
            changed = True
            logging.info("Boss reminder sent for %s records.", len(boss_records))

        if changed:
            wb.save(excel_path)
            logging.info("Workbook updated and saved.")
        else:
            logging.info("No reminder needed.")

        return 0
    except Exception:
        logging.exception("Runtime error.")
        logging.error(traceback.format_exc())
        return 1
    finally:
        if wb is not None:
            wb.close()


if __name__ == "__main__":
    raise SystemExit(main())
