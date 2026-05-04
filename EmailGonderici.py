#!/usr/bin/env python3
"""
Vehicle maintenance email notifier.

Config format (EmailGondericiAyarlar.txt):
- FILENAME: workbook_name.xlsx
- GMAILUSERNAME: your@gmail.com
- GMAILPASSWORD: app_password
- EMAILRECIPIENTS: a@x.com, b@y.com
- CALC: <plate_cell> <label_cell> <value_cell> <operator><threshold>
- SET: <cell> <increment_value>

Excel formulas in CALC value cells are evaluated in Python (no Excel install
required). Supported: =TODAY(), =NOW(), and binary arithmetic between two cell
references such as =D13-C13 or =(I22-H22).
"""

from __future__ import annotations

import argparse
import logging
import re
import smtplib
import sys
import traceback
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from email.message import EmailMessage
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import load_workbook


CONFIG_NAME = "EmailGondericiAyarlar.txt"
LOG_NAME = "EmailGonderici.log"
SMTP_HOST = "smtp.gmail.com"
SMTP_PORT = 465

CALC_PATTERN = re.compile(
    r"^(?P<plate>\S+)\s+(?P<label>\S+)\s+(?P<value>\S+)\s*(?P<op><=|>=|==|!=|<|>)\s*(?P<threshold>.+)$"
)
SET_PATTERN = re.compile(r"^(?P<cell>\S+)\s+(?P<increment>-?\d+(?:[.,]\d+)?)$")
CELL_REF_PATTERN = re.compile(r"^[A-Z]+\d+$")
BINARY_FORMULA_PATTERN = re.compile(
    r"^([A-Z]+\d+)\s*([\-+*/])\s*([A-Z]+\d+)$"
)


@dataclass
class CalcRule:
    plate_cell: str
    label_cell: str
    value_cell: str
    operator: str
    threshold: str


@dataclass
class SetRule:
    cell: str
    increment: float


@dataclass
class Config:
    filename: str = ""
    gmail_username: str = ""
    gmail_password: str = ""
    recipients: List[str] = field(default_factory=list)
    calc_rules: List[CalcRule] = field(default_factory=list)
    set_rules: List[SetRule] = field(default_factory=list)


def base_dir() -> Path:
    """Folder where config/log/excel live. Works for both .py and PyInstaller .exe."""
    if getattr(sys, "frozen", False):
        return Path(sys.executable).resolve().parent
    return Path(__file__).resolve().parent


def configure_logging(log_path: Path) -> None:
    logging.basicConfig(
        filename=str(log_path),
        level=logging.INFO,
        format="%(asctime)s - %(levelname)s - %(message)s",
    )


def read_lines_with_fallback(path: Path) -> List[str]:
    encodings = ("utf-8-sig", "utf-8", "cp1254", "latin-1")
    for encoding in encodings:
        try:
            return path.read_text(encoding=encoding).splitlines()
        except UnicodeDecodeError:
            continue
    raise UnicodeDecodeError("unknown", b"", 0, 1, "Unable to decode config file")


def parse_float(value: object) -> Optional[float]:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, timedelta):
        return float(value.days)
    text = str(value).strip()
    if not text:
        return None
    try:
        return float(text.replace(",", "."))
    except ValueError:
        return None


def parse_config(path: Path) -> Config:
    cfg = Config()
    lines = read_lines_with_fallback(path)

    for line_number, raw_line in enumerate(lines, start=1):
        line = raw_line.strip()
        if not line or line.startswith("#"):
            continue

        if ":" not in line:
            logging.warning("Line %s ignored (missing ':'): %s", line_number, raw_line)
            continue

        key, value = line.split(":", 1)
        key = key.strip().upper()
        value = value.strip()

        if key == "FILENAME":
            cfg.filename = value
            continue
        if key == "GMAILUSERNAME":
            cfg.gmail_username = value
            continue
        if key == "GMAILPASSWORD":
            cfg.gmail_password = value
            continue
        if key == "EMAILRECIPIENTS":
            cfg.recipients = [item.strip() for item in value.split(",") if item.strip()]
            continue
        if key == "CALC":
            match = CALC_PATTERN.match(value)
            if not match:
                logging.warning("Line %s invalid CALC rule: %s", line_number, raw_line)
                continue
            cfg.calc_rules.append(
                CalcRule(
                    plate_cell=match.group("plate"),
                    label_cell=match.group("label"),
                    value_cell=match.group("value"),
                    operator=match.group("op"),
                    threshold=match.group("threshold"),
                )
            )
            continue
        if key == "SET":
            match = SET_PATTERN.match(value)
            if not match:
                logging.warning("Line %s invalid SET rule: %s", line_number, raw_line)
                continue
            cfg.set_rules.append(
                SetRule(
                    cell=match.group("cell"),
                    increment=float(match.group("increment").replace(",", ".")),
                )
            )
            continue

        logging.warning("Line %s has unknown key '%s'.", line_number, key)

    required = {
        "FILENAME": cfg.filename,
        "GMAILUSERNAME": cfg.gmail_username,
        "GMAILPASSWORD": cfg.gmail_password,
        "EMAILRECIPIENTS": ",".join(cfg.recipients),
    }
    missing = [name for name, val in required.items() if not val]
    if missing:
        raise ValueError(f"Missing required config keys: {', '.join(missing)}")

    return cfg


def evaluate_condition(value: object, operator: str, threshold: str) -> bool:
    left_num = parse_float(value)
    right_num = parse_float(threshold)

    if left_num is not None and right_num is not None:
        left: object = left_num
        right: object = right_num
    else:
        left = str(value).strip()
        right = str(threshold).strip()
        if operator not in ("==", "!="):
            return False

    if operator == "<=":
        return left <= right  # type: ignore[operator]
    if operator == "<":
        return left < right  # type: ignore[operator]
    if operator == ">=":
        return left >= right  # type: ignore[operator]
    if operator == ">":
        return left > right  # type: ignore[operator]
    if operator == "==":
        return left == right
    if operator == "!=":
        return left != right
    return False


def format_cell_value(value: object) -> str:
    if value is None:
        return "-"
    if isinstance(value, datetime):
        return value.strftime("%d.%m.%Y")
    if isinstance(value, date):
        return value.strftime("%d.%m.%Y")
    if isinstance(value, timedelta):
        return str(value.days)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def to_date(value: object) -> Optional[date]:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return None


def evaluate_binary(left: object, op: str, right: object) -> object:
    left_date = to_date(left)
    right_date = to_date(right)

    if left_date is not None and right_date is not None:
        if op == "-":
            return (left_date - right_date).days
        if op == "+":
            raise ValueError("Date addition not supported.")

    left_num = parse_float(left)
    right_num = parse_float(right)
    if left_num is None or right_num is None:
        raise ValueError(f"Non-numeric operands: {left!r} {op} {right!r}")

    if op == "-":
        return left_num - right_num
    if op == "+":
        return left_num + right_num
    if op == "*":
        return left_num * right_num
    if op == "/":
        if right_num == 0:
            raise ZeroDivisionError("Division by zero in formula.")
        return left_num / right_num
    raise ValueError(f"Unsupported operator: {op}")


class FormulaEvaluator:
    """Evaluate the small set of Excel formulas used in our workbooks."""

    def __init__(self, ws: Any, today_value: date) -> None:
        self.ws = ws
        self.today = today_value
        self.cache: Dict[str, object] = {}
        self.in_progress: set = set()

    def value(self, ref: str) -> object:
        ref_upper = ref.upper()
        if ref_upper in self.cache:
            return self.cache[ref_upper]
        if ref_upper in self.in_progress:
            raise ValueError(f"Circular reference at {ref_upper}")

        cell = self.ws[ref_upper]
        raw = cell.value

        if isinstance(raw, str) and raw.startswith("="):
            self.in_progress.add(ref_upper)
            try:
                result = self._eval_formula(raw)
            finally:
                self.in_progress.discard(ref_upper)
        else:
            result = raw

        self.cache[ref_upper] = result
        return result

    def _eval_formula(self, formula: str) -> object:
        expr = formula.strip()
        if expr.startswith("="):
            expr = expr[1:].strip()

        upper = expr.upper().replace(" ", "")
        if upper in ("TODAY()", "NOW()"):
            return self.today

        if expr.startswith("(") and expr.endswith(")"):
            expr = expr[1:-1].strip()

        if CELL_REF_PATTERN.match(expr.upper()):
            return self.value(expr)

        match = BINARY_FORMULA_PATTERN.match(expr.upper())
        if match:
            ref1, op, ref2 = match.group(1), match.group(2), match.group(3)
            left = self.value(ref1)
            right = self.value(ref2)
            return evaluate_binary(left, op, right)

        try:
            return float(expr)
        except ValueError:
            pass

        raise ValueError(f"Unsupported formula: {formula!r}")


def apply_set_rules(ws: Any, set_rules: List[SetRule]) -> None:
    for rule in set_rules:
        current = ws[rule.cell].value
        current_num = parse_float(current)
        if current_num is None:
            logging.warning("SET skipped, non-numeric value at %s: %r", rule.cell, current)
            continue
        new_value = current_num + rule.increment
        if float(new_value).is_integer():
            new_value = int(new_value)
        ws[rule.cell].value = new_value


def evaluate_calc_rules(
    ws: Any, calc_rules: List[CalcRule], today_value: date
) -> List[Tuple[str, str, str, str]]:
    evaluator = FormulaEvaluator(ws, today_value)
    matched: List[Tuple[str, str, str, str]] = []

    for rule in calc_rules:
        try:
            plate = evaluator.value(rule.plate_cell)
            label = evaluator.value(rule.label_cell)
            value = evaluator.value(rule.value_cell)
        except Exception:
            logging.exception(
                "CALC kuralinda hucre okuma hatasi: %s %s %s",
                rule.plate_cell,
                rule.label_cell,
                rule.value_cell,
            )
            continue

        if evaluate_condition(value, rule.operator, rule.threshold):
            matched.append(
                (
                    format_cell_value(plate),
                    format_cell_value(label),
                    format_cell_value(value),
                    f"{rule.operator}{rule.threshold}",
                )
            )

    return matched


def process_workbook(
    workbook_path: Path,
    cfg: Config,
    apply_set: bool = True,
    today_value: Optional[date] = None,
) -> List[Tuple[str, str, str, str]]:
    if today_value is None:
        today_value = date.today()

    if apply_set and cfg.set_rules:
        wb = load_workbook(filename=str(workbook_path), data_only=False)
        ws = wb.worksheets[0]
        apply_set_rules(ws, cfg.set_rules)
        wb.save(str(workbook_path))
        wb.close()
        logging.info("SET kurallari kaydedildi (%s kural).", len(cfg.set_rules))
    else:
        logging.info("SET kurallari atlandi (dry-run veya bos liste).")

    wb_calc = load_workbook(filename=str(workbook_path), data_only=False)
    ws_calc = wb_calc.worksheets[0]
    try:
        matched = evaluate_calc_rules(ws_calc, cfg.calc_rules, today_value)
        logging.info(
            "CALC degerlendirildi: %s kural, %s eslesme.",
            len(cfg.calc_rules),
            len(matched),
        )
        return matched
    finally:
        wb_calc.close()


def build_email_body(records: List[Tuple[str, str, str, str]]) -> str:
    lines = [
        "Merhaba,",
        "",
        "Asagidaki arac bakim kalemleri icin esik degeri asildi:",
        "",
    ]
    for plate, label, value, _condition in records:
        lines.append(f"- Plaka: {plate} | Kalem: {label} | Deger: {value}")
    lines.extend(["", "Bu e-posta otomatik olarak gonderilmistir."])
    return "\n".join(lines)


def send_email(cfg: Config, records: List[Tuple[str, str, str, str]]) -> None:
    if not cfg.recipients:
        raise ValueError("EMAILRECIPIENTS bos olamaz.")

    message = EmailMessage()
    message["Subject"] = f"Arac Bakim Uyarisi - {datetime.now():%Y-%m-%d %H:%M}"
    message["From"] = cfg.gmail_username
    message["To"] = ", ".join(cfg.recipients)
    message.set_content(build_email_body(records))

    with smtplib.SMTP_SSL(SMTP_HOST, SMTP_PORT, timeout=30) as smtp:
        smtp.login(cfg.gmail_username, cfg.gmail_password)
        smtp.send_message(message)


def ensure_config(script_dir: Path, force_setup: bool) -> bool:
    """Open setup wizard if config missing or --setup requested. Returns True if continuing."""
    config_path = script_dir / CONFIG_NAME
    if not force_setup and config_path.exists():
        return True

    try:
        from setup_wizard import run_arac_takip_wizard
    except ImportError:
        if not config_path.exists():
            logging.error(
                "Config dosyasi yok ve setup wizard yuklenemedi. %s olusturup tekrar deneyin.",
                config_path,
            )
            return False
        return True

    return run_arac_takip_wizard(config_path)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="Excel kurallarini calistirir, e-posta gondermez.",
    )
    parser.add_argument(
        "--setup",
        action="store_true",
        help="Ayar sihirbazini ac (Tkinter formu).",
    )
    args = parser.parse_args()

    script_dir = base_dir()
    log_path = script_dir / LOG_NAME
    configure_logging(log_path)

    if not ensure_config(script_dir, force_setup=args.setup):
        return 1

    config_path = script_dir / CONFIG_NAME
    if not config_path.exists():
        logging.error("Config file not found: %s", config_path)
        return 1

    try:
        cfg = parse_config(config_path)
        logging.info("%s Ayar dosyasi okundu.", CONFIG_NAME)
    except Exception:
        logging.exception("Ayar dosyasini okurken hata olustu. Detaylari altta.")
        logging.error(traceback.format_exc())
        return 1

    workbook_path = script_dir / cfg.filename
    if not workbook_path.exists():
        logging.error("Excel dosyasi bulunamadi: %s", workbook_path)
        return 1

    try:
        matched = process_workbook(workbook_path, cfg, apply_set=not args.dry_run)
    except Exception:
        logging.exception("Excel islemleri ile ilgili hata oldu. Detaylari altta.")
        logging.error(traceback.format_exc())
        return 1

    try:
        if not matched:
            logging.info("No available records exist to send email.")
            if args.dry_run:
                print("DRY RUN: no records matched")
            return 0

        if args.dry_run:
            print(f"DRY RUN: matched {len(matched)} records")
            for plate, label, value, condition in matched[:10]:
                print(f"- {plate} | {label} | {value} | {condition}")
            if len(matched) > 10:
                print(f"... and {len(matched) - 10} more")
            logging.info("Dry run mode active, email gonderimi atlandi.")
            return 0

        send_email(cfg, matched)
        logging.info("Email Sent")
        return 0
    except Exception:
        logging.exception("Email gonderimi ile ilgili hata oldu. Detaylari altta.")
        logging.error(traceback.format_exc())
        return 1
    finally:
        logging.info("Email session closed.")


if __name__ == "__main__":
    raise SystemExit(main())
